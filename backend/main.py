import io
import pdfplumber
from docx import Document
from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse

from models import ResumeData
from ai_service import (
    optimize_resume_with_ai,
    parse_resume_from_text,
    calculate_ats_score,
    tailor_skills_with_ai
)
from latex_export import (
    generate_resume_latex,
    generate_resume_latex_zip,
    compile_latex_to_pdf_with_verification
)
from docx_export import generate_resume_docx

app = FastAPI(title="ATSForge API", description="AI-powered ATS Resume Builder & Optimizer API")

# Enable CORS for the React frontend (Vite defaults to 5173 / 3000)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/")
def read_root():
    return {"message": "Welcome to the ATSForge API"}

@app.get("/health")
def health_check():
    return {"status": "ok", "service": "ATSForge API"}

@app.post("/api/optimize")
def optimize_resume(resume: ResumeData):
    optimized_resume = optimize_resume_with_ai(resume)
    return {"status": "success", "message": "Resume optimized successfully.", "data": optimized_resume}

@app.post("/api/export-docx")
def export_docx(resume: ResumeData):
    file_stream = generate_resume_docx(resume)
    return StreamingResponse(
        file_stream, 
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": "attachment; filename=Resume_ATS_Export.docx"
        }
    )

@app.post("/api/upload-resume")
async def upload_resume(file: UploadFile = File(...)):
    raw_text = ""
    contents = await file.read()
    
    if file.filename.lower().endswith(".pdf"):
        try:
            with pdfplumber.open(io.BytesIO(contents)) as pdf:
                for page in pdf.pages:
                    # Attempt layout=True first to preserve column separation in multi-column resumes
                    page_text = page.extract_text(layout=True)
                    if not page_text or len(page_text.strip()) < 40:
                        # Fallback to standard x/y tolerance extraction
                        page_text = page.extract_text(x_tolerance=2, y_tolerance=2)
                    if page_text:
                        raw_text += page_text + "\n\n"
        except Exception as e:
            print(f"PDF Extract Error: {e}")
    elif file.filename.lower().endswith(".docx"):
        try:
            doc = Document(io.BytesIO(contents))
            parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        # Deduplicate adjacent cells caused by merged table cells in Word
                        parts.append(" | ".join(dict.fromkeys(cells)))
            raw_text = "\n".join(parts)
        except Exception as e:
            print(f"DOCX Extract Error: {e}")
            
    if not raw_text.strip():
        return {"status": "error", "message": "Could not extract text from file."}
        
    parsed_data = parse_resume_from_text(raw_text)
    if not parsed_data:
        return {"status": "error", "message": "AI could not parse the resume. Please check your API key or document format."}
    return {"status": "success", "data": parsed_data}

@app.post("/api/ats-score")
def ats_score(resume: ResumeData):
    result = calculate_ats_score(resume)
    return {"status": "success", "data": result}

@app.post("/api/tailor-skills")
def tailor_skills(resume: ResumeData):
    result = tailor_skills_with_ai(resume)
    return result

@app.post("/api/export-latex")
def export_latex(resume: ResumeData):
    latex_code = generate_resume_latex(resume, bullet_budget=3, spacing_level="auto")
    return {"status": "success", "latex": latex_code}

@app.post("/api/export-latex-zip")
def export_latex_zip(resume: ResumeData):
    zip_buffer = generate_resume_latex_zip(resume)
    filename = f"{resume.full_name.replace(' ', '_') or 'Resume'}_Kyvernitis_LaTeX.zip"
    return StreamingResponse(
        zip_buffer,
        media_type="application/zip",
        headers={
            "Content-Disposition": f"attachment; filename={filename}"
        }
    )

@app.post("/api/compile-latex")
def compile_latex(resume: ResumeData):
    success, pages, pdf_bytes, latex_code = compile_latex_to_pdf_with_verification(resume)
    if success and pdf_bytes:
        filename = f"{resume.full_name.replace(' ', '_') or 'Resume'}_ATS_Kyvernitis.pdf"
        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={
                "Content-Disposition": f"attachment; filename={filename}"
            }
        )
    return {
        "status": "partial_success",
        "verified_pages": pages,
        "compiler_available": False if not pdf_bytes else True,
        "message": "LaTeX generated and pre-budgeted for 1-page layout. To compile to PDF on your machine, install XeLaTeX, or use the Download LaTeX Zip button to open in Overleaf.",
        "latex": latex_code
    }
